#!/usr/bin/env python3
"""
PDF en Excel tekst extractor voor Ziraat raporlari.
Gebruik: python3 lees_pdf.py "pad/naar/bestand.pdf"
         python3 lees_pdf.py "pad/naar/bestand.xlsx"
         python3 lees_pdf.py --lijst "pad/naar/map"
         python3 lees_pdf.py --recent [aantal] "pad/naar/map"
"""

import sys
import os
import fitz  # PyMuPDF
import openpyxl


def lees_pdf(pad):
    try:
        with fitz.open(pad) as doc:
            tekst = "".join(pagina.get_text() for pagina in doc)
        return tekst.strip()
    except Exception as e:
        return f"HATA: {pad} okunamadi: {e}"


def lees_excel(pad):
    try:
        # read_only: streamt rijen i.p.v. het hele werkboek in geheugen te laden
        wb = openpyxl.load_workbook(pad, data_only=True, read_only=True)
        try:
            regels = []
            for bladnaam in wb.sheetnames:
                blad = wb[bladnaam]
                regels.append(f"=== Blad: {bladnaam} ===")
                for rij in blad.iter_rows(values_only=True):
                    if any(cel is not None for cel in rij):
                        regels.append("\t".join("" if cel is None else str(cel) for cel in rij))
        finally:
            wb.close()
        return "\n".join(regels).strip()
    except Exception as e:
        return f"HATA: {pad} okunamadi: {e}"


def lees_bestand(pad):
    if pad.lower().endswith(".xlsx"):
        return lees_excel(pad)
    return lees_pdf(pad)


def lijst_bestanden(map_pad):
    resultaten = []
    for root, dirs, files in os.walk(map_pad):
        dirs.sort()  # deterministische volgorde over mappen heen
        for bestand in sorted(files):
            if bestand.lower().endswith((".pdf", ".xlsx")):
                resultaten.append(os.path.join(root, bestand))
    return resultaten


def recente_bestanden(map_pad, aantal):
    bestanden = lijst_bestanden(map_pad)
    bestanden.sort(key=os.path.getmtime, reverse=True)
    return bestanden[:aantal]


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Gebruik: python3 lees_pdf.py <bestand.pdf|bestand.xlsx>")
        print("         python3 lees_pdf.py --lijst <map>")
        print("         python3 lees_pdf.py --recent [aantal] <map>")
        sys.exit(1)

    if sys.argv[1] == "--lijst":
        map_pad = sys.argv[2] if len(sys.argv) > 2 else "."
        for p in lijst_bestanden(map_pad):
            print(p)
    elif sys.argv[1] == "--recent":
        rest = sys.argv[2:]
        aantal = 15
        if rest and rest[0].isdigit():
            aantal = int(rest[0])
            rest = rest[1:]
        map_pad = rest[0] if rest else "."
        for p in recente_bestanden(map_pad, aantal):
            print(p)
    else:
        pad = sys.argv[1]
        uit = lees_bestand(pad)
        print(uit)
        if uit.startswith("HATA:"):
            sys.exit(1)
        if pad.lower().endswith(".pdf") and len(uit) < 500:
            print(f"LET OP: slechts {len(uit)} tekens geëxtraheerd; "
                  "PDF is waarschijnlijk image-based.", file=sys.stderr)
